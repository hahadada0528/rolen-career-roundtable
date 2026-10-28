"""作品集导出：与 GEO 工具同一套 reportlab / python-docx 方法。"""
from __future__ import annotations

import re
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EXPORT_DIR = ROOT / "outputs"


def _safe_stem(text: str, fallback: str = "portfolio") -> str:
    raw = (text or "").strip() or fallback
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", raw)
    cleaned = cleaned.strip(" ._")[:60] or fallback
    return cleaned


def _split_blocks(article: str) -> list[tuple[str, str]]:
    """粗分标题/段落。返回 [(kind, text)] kind in heading1|heading2|para。"""
    blocks: list[tuple[str, str]] = []
    lines = (article or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            blocks.append(("heading1", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("heading2", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("heading2", line[4:].strip()))
        elif re.match(r"^【.+】$", line) or re.match(r"^第[一二三四五六七八九十\d]+[章节部分]", line):
            blocks.append(("heading2", line))
        else:
            blocks.append(("para", line))
    if not blocks:
        blocks.append(("para", (article or "").strip() or "（空文稿）"))
    if blocks and all(k == "para" for k, _ in blocks):
        first = blocks[0][1]
        if len(first) <= 40:
            blocks[0] = ("heading1", first)
    return blocks


def _pick_chinese_font() -> str | None:
    candidates = [
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    ]
    for p in candidates:
        if p.exists():
            return str(p)
    return None


def export_pdf(article: str, *, title: str) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = EXPORT_DIR / f"{_safe_stem(title)}_{stamp}.pdf"

    font_path = _pick_chinese_font()
    font_name = "Helvetica"
    if font_path:
        try:
            pdfmetrics.registerFont(TTFont("GeoCN", font_path))
            font_name = "GeoCN"
        except Exception:
            try:
                pdfmetrics.registerFont(TTFont("GeoCN", font_path, subfontIndex=0))
                font_name = "GeoCN"
            except Exception:
                font_name = "Helvetica"

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "GeoBody",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=11,
        leading=18,
        spaceAfter=10,
    )
    h1 = ParagraphStyle(
        "GeoH1",
        parent=styles["Heading1"],
        fontName=font_name,
        fontSize=16,
        leading=22,
        spaceAfter=12,
        spaceBefore=6,
    )
    h2 = ParagraphStyle(
        "GeoH2",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=13,
        leading=20,
        spaceAfter=8,
        spaceBefore=8,
    )
    meta = ParagraphStyle(
        "GeoMeta",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=10,
        leading=15,
        textColor="#445566",
        spaceAfter=14,
    )

    def esc(t: str) -> str:
        return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=title or "portfolio",
    )
    story = []
    first = True
    for kind, text in _split_blocks(article):
        if kind == "heading1":
            story.append(Paragraph(esc(text), h1))
            if first:
                story.append(
                    Paragraph(
                        esc(
                            f"作品集案例 · 导出时间 {datetime.now().strftime('%Y-%m-%d')} · "
                            "https://github.com/hahadada0528/rolen-career-roundtable"
                        ),
                        meta,
                    )
                )
                first = False
        elif kind == "heading2":
            story.append(Paragraph(esc(text), h2))
        else:
            story.append(Paragraph(esc(text), body))
    story.append(Spacer(1, 12))
    doc.build(story)
    return path


def export_docx(article: str, *, title: str) -> Path:
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = EXPORT_DIR / f"{_safe_stem(title)}_{stamp}.docx"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for kind, text in _split_blocks(article):
        if kind == "heading1":
            doc.add_heading(text, level=1)
        elif kind == "heading2":
            doc.add_heading(text, level=2)
        else:
            doc.add_paragraph(text)

    doc.save(str(path))
    return path


def main() -> int:
    md_path = ROOT / "portfolio" / "作品集-睿览职业工作手册圆桌.md"
    if not md_path.is_file():
        print(f"缺少文稿：{md_path}", file=sys.stderr)
        return 1
    article = md_path.read_text(encoding="utf-8")
    title = "睿览·职业工作手册圆桌（作品集）"
    for line in article.splitlines():
        if line.startswith("# "):
            title = line[2:].strip()
            break
    pdf = export_pdf(article, title=title)
    print(f"PDF: {pdf}")
    try:
        docx = export_docx(article, title=title)
        print(f"DOCX: {docx}")
    except Exception as exc:
        print(f"DOCX 跳过：{exc}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
